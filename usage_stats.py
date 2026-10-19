# Local usage statistics - no message bodies or passwords
import csv
import sqlite3
from contextlib import contextmanager
import time
from datetime import datetime, timezone
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from mail_backend import clean

@contextmanager
def open_db(path):
    """Unlike sqlite3's own context manager, this ALSO closes connections."""
    con = sqlite3.connect(path, timeout=30)
    try:
        con.execute('PRAGMA busy_timeout=30000')
        yield con
        con.commit()
    except Exception:
        con.rollback()
        raise
    finally:
        con.close()

COLUMNS = ['id','account','started_at','finished_at','status','saved','skipped','errors','download_bytes','storage_bytes','seconds','average_mb_s']

def root_db(destination):
    root = Path(destination) / '_BAO_CAO_SU_DUNG'
    root.mkdir(parents=True, exist_ok=True)
    return root

class UsageDB:
    def __init__(self, destination):
        self.path = root_db(destination) / 'lich_su.sqlite3'
        with open_db(self.path) as con:
            con.execute('''CREATE TABLE IF NOT EXISTS runs (id INTEGER PRIMARY KEY AUTOINCREMENT, account TEXT NOT NULL,
            started_at TEXT, finished_at TEXT, status TEXT, saved INTEGER DEFAULT 0, skipped INTEGER DEFAULT 0,
            errors INTEGER DEFAULT 0, download_bytes INTEGER DEFAULT 0, storage_bytes INTEGER DEFAULT 0,
            seconds REAL DEFAULT 0, average_mb_s REAL DEFAULT 0)''')

    def start(self, account):
        now = datetime.now(timezone.utc).isoformat()
        with open_db(self.path) as con:
            cur = con.execute('INSERT INTO runs(account,started_at,status) VALUES (?,?,?)', (account,now,'running'))
            return cur.lastrowid

    def finish(self, session_id, stats, status, storage_bytes):
        now = datetime.now(timezone.utc)
        with open_db(self.path) as con:
            start = con.execute('SELECT started_at FROM runs WHERE id=?',(session_id,)).fetchone()
            if not start: return
            secs = max((now - datetime.fromisoformat(start[0])).total_seconds(),.001)
            bytecount = stats.get('bytes',0)
            con.execute('''UPDATE runs SET finished_at=?,status=?,saved=?,skipped=?,errors=?,
            download_bytes=?,storage_bytes=?,seconds=?,average_mb_s=? WHERE id=?''',
            (now.isoformat(),status,stats.get('saved',0),stats.get('skipped',0),stats.get('errors',0),
             bytecount,storage_bytes,secs,bytecount/secs/1048576,session_id))

def storage_usage(destination, account):
    folder = Path(destination) / clean(account)
    if not folder.exists(): return 0
    total = 0
    for p in folder.rglob('*'):
        try:
            if p.is_file(): total += p.stat().st_size
        except OSError: pass
    return total

def spreadsheet_safe(value):
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@', '\t', '\r')):
        return "'" + value
    return value

def export_reports(destination):
    directory = root_db(destination)
    dbpath = directory/'lich_su.sqlite3'
    if not dbpath.exists(): UsageDB(destination)
    with open_db(dbpath) as con:
        records = con.execute('SELECT '+','.join(COLUMNS)+' FROM runs ORDER BY id').fetchall()
    csvpath=directory/'lich_su_sao_luu.csv'
    with csvpath.open('w', encoding='utf-8-sig', newline='') as f:
        w=csv.writer(f); w.writerow(COLUMNS); w.writerows([[spreadsheet_safe(v) for v in row] for row in records])
    summary={}
    for row in records:
        acct=row[1]; t=summary.setdefault(acct,[0,0,0,0,0,0])
        t[0]+=1; t[1]+=row[5]; t[2]+=row[6]; t[3]+=row[7]; t[4]+=row[8]; t[5]=max(t[5],row[9])
    with (directory/'tong_hop_nhan_vien.csv').open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f); w.writerow(['account','runs','saved','skipped','errors','download_MB','storage_MB'])
        for a, t in sorted(summary.items()): w.writerow([spreadsheet_safe(a),*t[:4],round(t[4]/1048576,2),round(t[5]/1048576,2)])
    book=Workbook(); ws=book.active; ws.title='Lich su'; ws.append(COLUMNS)
    for rec in records: ws.append([spreadsheet_safe(v) for v in rec])
    summ=book.create_sheet('Tong hop'); summ.append(['Tai khoan','So luot','Da luu','Bo qua','Loi','Da tai (MB)','Dung luong luu (MB)'])
    for a,t in sorted(summary.items()): summ.append([spreadsheet_safe(a),*t[:4],round(t[4]/1048576,2),round(t[5]/1048576,2)])
    for sh in book:
        sh.freeze_panes='A2'; sh.auto_filter.ref=sh.dimensions
        for cell in sh[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='245B84')
        for col in sh.columns:
            letter=col[0].column_letter
            sh.column_dimensions[letter].width=min(35,max(13,max(len(str(c.value or '')) for c in col)+2))
    book.save(directory/'bao_cao_sao_luu.xlsx')
    return directory
